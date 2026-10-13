from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_title_slide(prs, title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    # Set dark background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(17, 24, 39) # Dark Blue/Grey

    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(59, 130, 246) # Blue Accent
    title.text_frame.paragraphs[0].font.size = Pt(54)
    title.text_frame.paragraphs[0].font.bold = True

    subtitle = slide.placeholders[1]
    subtitle.text = subtitle_text
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(209, 213, 219)
    subtitle.text_frame.paragraphs[0].font.size = Pt(24)

def add_content_slide(prs, title_text, bullets, accent_color):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(249, 250, 251) # Very light grey

    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = accent_color
    title.text_frame.paragraphs[0].font.bold = True

    content = slide.placeholders[1].text_frame
    content.clear()
    
    for idx, bullet in enumerate(bullets):
        p = content.add_paragraph()
        if isinstance(bullet, tuple): # It's a sub-bullet
            p.text = bullet[0]
            p.level = 1
            p.font.size = Pt(20)
            p.font.color.rgb = RGBColor(75, 85, 99)
        else:
            p.text = bullet
            p.level = 0
            p.font.size = Pt(24)
            p.font.color.rgb = RGBColor(31, 41, 55)
            # Add a bit of spacing before main bullets
            if idx > 0:
                p.space_before = Pt(14)

def create_hackathon_ppt():
    prs = Presentation()
    
    # 1. Title Slide
    add_title_slide(prs, "AI Smart City Dashboard", "Intelligent Urban Management & Predictive Analytics")

    # 2. Problem Statement
    add_content_slide(prs, "The Problem", [
        "Rapid Urbanization & Inefficiency",
        ("Traffic congestion leads to immense time and fuel waste.",),
        ("Waste collection schedules are static, leading to overflowing bins.",),
        ("Citizens lack a centralized platform to report localized issues.",),
        "Data Silos in City Management",
        ("Departments (Traffic, Environment, Sanitation) operate independently.",),
        ("Lack of predictive insights to proactively solve problems.",)
    ], RGBColor(220, 38, 38)) # Red

    # 3. Our Solution
    add_content_slide(prs, "Our Solution", [
        "A Centralized, AI-Driven Platform",
        "Real-Time Analytics Dashboard",
        ("Live monitoring of traffic, AQI, noise, and waste fill levels.",),
        "Machine Learning Predictions",
        ("Forecasting congestion 30 mins ahead and bin fill levels 5 hours ahead.",),
        "Interactive Citizen Portal",
        ("Empowering citizens to report issues and access a smart AI assistant.",)
    ], RGBColor(16, 185, 129)) # Green

    # 4. Key Features Deep-Dive
    add_content_slide(prs, "Key Features", [
        "🚦 Intelligent Traffic Management",
        ("Live Leaflet.js mapping with congestion tooltips.",),
        ("Random Forest regression for highly accurate congestion forecasting.",),
        "♻️ Smart Waste Management",
        ("Linear regression to predict future bin fill percentages.",),
        "🗣️ AI Chatbot Assistant",
        ("Answers natural language queries regarding city data & alerts.",),
        "📝 Automated Complaint Routing",
        ("Citizens report issues directly tied to specific city zones.",)
    ], RGBColor(139, 92, 246)) # Purple

    # 5. Tech Stack
    add_content_slide(prs, "Technology Stack", [
        "Backend / API",
        ("Python, Flask, SQLite, SQLAlchemy",),
        "Machine Learning & Data",
        ("Scikit-Learn (RandomForestRegressor, LinearRegression)",),
        ("APScheduler for real-time IoT sensor simulation",),
        "Frontend Interface",
        ("Vanilla JavaScript (ES6+), Fetch API, HTML5/CSS3",),
        ("Leaflet.js for interactive mapping",)
    ], RGBColor(59, 130, 246)) # Blue

    # 6. Future Scope
    add_content_slide(prs, "Future Scope & Scalability", [
        "Infrastructure Upgrades",
        ("Migrate from SQLite to PostgreSQL for high concurrency.",),
        ("Implement WebSockets (Socket.IO) to replace HTTP polling.",),
        "Advanced AI Enhancements",
        ("Deep Learning (LSTMs) for complex temporal traffic patterns.",),
        ("Computer Vision integration for automated pothole detection via CCTVs.",),
        "Mobile App Integration",
        ("Push notifications for critical city alerts and AQI drops.",)
    ], RGBColor(245, 158, 11)) # Orange

    # 7. Thank you
    add_title_slide(prs, "Thank You", "Ready for Q&A!")

    prs.save("Hackathon_Pitch_Deck.pptx")

create_hackathon_ppt()
