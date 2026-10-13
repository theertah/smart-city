import collections
import collections.abc
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_ppt():
    prs = pptx.Presentation()
    
    # Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AI Smart City Platform"
    subtitle.text = "Intelligent Urban Management Dashboard\nCreated: 2026"
    
    # Make title colorful
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(37, 99, 235) # Accent Blue
    
    # Agenda Slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Project Overview"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(16, 185, 129) # Success Green
    
    content = slide.placeholders[1].text_frame
    content.text = "1. Introduction to Smart City Dashboard"
    p = content.add_paragraph()
    p.text = "2. Live Traffic & Congestion Management"
    p = content.add_paragraph()
    p.text = "3. Environmental & Weather Monitoring"
    p = content.add_paragraph()
    p.text = "4. AI Waste Management System"
    p = content.add_paragraph()
    p.text = "5. Citizen Portal & Complaints"
    p = content.add_paragraph()
    p.text = "6. Smart Assistant Integration"
    
    # Traffic Slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Live Traffic & Predictions"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(239, 68, 68) # Danger Red
    content = slide.placeholders[1].text_frame
    content.text = "• Interactive Map (Bengaluru)"
    p = content.add_paragraph()
    p.text = "• Simulated Live Data (Speed, Count, Congestion %)"
    p = content.add_paragraph()
    p.text = "• Predictive AI Models (30-min congestion forecasting)"
    p = content.add_paragraph()
    p.text = "• Real-time Zone Metrics"

    # Environment Slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Environment Monitoring"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(245, 158, 11) # Warning Yellow
    content = slide.placeholders[1].text_frame
    content.text = "• Live Weather Integration (Temperature, Wind, Humidity)"
    p = content.add_paragraph()
    p.text = "• Air Quality Index (AQI) tracking"
    p = content.add_paragraph()
    p.text = "• Noise level monitoring across zones"
    
    # Citizen & Assistant Slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Citizen Portal & Smart Assistant"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(139, 92, 246) # Purple
    content = slide.placeholders[1].text_frame
    content.text = "• File & Track City Complaints"
    p = content.add_paragraph()
    p.text = "• Dark / Light Theme customization"
    p = content.add_paragraph()
    p.text = "• AI Chatbot with Google Search grounding"
    p = content.add_paragraph()
    p.text = "• Interactive natural language Q&A about the city"

    prs.save("Smart_City_Presentation.pptx")

create_ppt()
